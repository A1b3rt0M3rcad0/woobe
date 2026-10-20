import hashlib
import importlib
import json
from uuid import uuid4

import pytest
from sqlalchemy import func, select, text
from sqlalchemy.exc import DBAPIError

from modules.artifacts.background_boundary import build_processing_runner
from modules.artifacts.composition import processing
from modules.artifacts.infra.database.models import ProcessingJobModel, RevisionModel, WorkspaceModel
from modules.artifacts.infra.database.representation_models import ArtifactRepresentationSetModel as Generation, ArtifactStoredObjectModel as StoredObject
from modules.artifacts.infra.database.unit_of_work import ArtifactsUnitOfWork
from tests.integration.artifacts.test_http_postgres import http_client as http_client, provision
from tests.integration.artifacts.test_profiles_postgres import create_profile
from tests.integration.artifacts.test_uploads_postgres import object_store as object_store
from woobe_workers.settings import WorkerSettings


async def prepared(http_client, engine, store, monkeypatch, *, quota=1_048_576):
    client, app, definition, actor, control, control_id = http_client
    for name in ('get_representation', 'list_representation_entries', 'resolve_source'):
        module = importlib.import_module('modules.artifacts.composition.composers.' + name)
        monkeypatch.setattr(module, 'DATABASE_ENGINE', engine)
    module = importlib.import_module('modules.artifacts.composition.representations')
    monkeypatch.setattr(module, 'upload_storage', lambda profile: store)
    monkeypatch.setattr(processing, 'upload_storage', lambda profile: store)
    workspace = await provision(client, definition.model_copy(update={'quota_bytes': quota}))
    wid = workspace['id']
    original = 'Primeiro parágrafo com ação e 😀.\n\nSegundo parágrafo.\n'
    response = await client.post(f'/v1/artifact-workspaces/{wid}/artifacts', json={
        'operation_id': uuid4().hex, 'kind': 'text', 'path': 'source.txt', 'text': original})
    assert response.status_code == 201, response.text
    ref = response.json()['receipt']
    profile = await create_profile(client, control, recipe={
        'mode': 'extract', 'processor': {'name': 'native-text', 'version': '1'}})
    assert profile.status_code == 201, profile.text
    body = {'operation_id': uuid4().hex, 'revision_id': ref['revision_id'],
            'profile_id': profile.json()['id'], 'profile_revision': 1}
    params = {'artifact_workspace_id': wid}
    path = '/v1/artifacts/' + ref['artifact_id'] + '/processing'
    response = await client.post(path, json=body, params=params)
    assert response.status_code == 202, response.text
    return client, wid, ref, original, path, body, response.json(), params


async def completed(client, path, job, params):
    response = await client.get(path + '/' + job['id'], params=params)
    assert response.status_code == 200, response.text
    return response.json()


async def test_atomic_bundle_sources_readonly_download_editable_copy_and_scoped_reuse(http_client, engine, object_store, monkeypatch):
    client, wid, ref, original, path, body, job, params = await prepared(http_client, engine, object_store, monkeypatch)
    base = f"/v1/artifacts/{ref['artifact_id']}/revisions/{ref['revision_id']}/representations"
    assert (await client.get(base, params=params)).status_code == 409
    await build_processing_runner(engine, WorkerSettings.model_validate({})).run_once('bundle-native-worker')
    job = await completed(client, path, job, params)
    assert job['status'] == 'succeeded', job
    generation = job['result']['representation_set_id']
    response = await client.get(base, params=params)
    assert response.status_code == 200, response.text
    assert response.json()['id'] == generation
    assert response.json()['source_sha256'] == hashlib.sha256(original.encode()).hexdigest()
    assert 'object_key' not in response.text and 'storage_version' not in response.text
    root = base + '/' + generation
    entries, cursor = [], None
    while True:
        response = await client.get(root + '/entries', params={**params, 'limit': 1, **({'cursor': cursor} if cursor else {})})
        assert response.status_code == 200, response.text
        page = response.json()
        entries.extend(page['items'])
        if not page['has_more']:
            break
        cursor = page['next_cursor']
    assert {entry['kind'] for entry in entries} == {'markdown', 'structure', 'source_map', 'outline'}
    assert all(entry['read_only'] for entry in entries)
    assert len({entry['id'] for entry in entries}) == len(entries)
    source_entry = next(entry for entry in entries if entry['kind'] == 'source_map')
    source_map = await client.get(root + '/entries/' + source_entry['id'] + '/content', params=params)
    assert source_map.status_code == 200, source_map.text
    source = source_map.json()['sources'][0]['source']
    read = await client.post('/v1/artifact-sources/read', params=params, json={'source': source, 'max_chars': 8})
    assert read.status_code == 200, read.text
    assert read.json()['text'] == 'Primeiro' and read.json()['truncated']
    assert read.json()['source']['representation_set_id'] == generation
    all_text = await client.post('/v1/artifact-sources/read', params=params, json={'source': source})
    assembled, offset = '', 0
    while True:
        part = await client.post('/v1/artifact-sources/read', params=params,
            json={'source': source, 'offset': offset, 'max_chars': 7})
        assert part.status_code == 200, part.text
        assembled += part.json()['text']
        if part.json()['next_offset'] is None:
            break
        offset = part.json()['next_offset']
    assert assembled == all_text.json()['text'] and '😀' in assembled
    for field, value in [('excerpt_sha256', '0' * 64), ('revision_id', 'missing-revision')]:
        invalid = {**source, field: value}
        assert (await client.post('/v1/artifact-sources/read', params=params, json={'source': invalid})).status_code == 404
    invalid = {**source, 'locator': {**source['locator'], 'block_id': 'missing-block'}}
    assert (await client.post('/v1/artifact-sources/read', params=params, json={'source': invalid})).status_code == 404
    bad = json.loads(json.dumps(source))
    bad['source_sha256'] = '0' * 64
    assert (await client.post('/v1/artifact-sources/read', params=params, json={'source': bad})).status_code == 404
    bad = json.loads(json.dumps(source))
    bad['locator']['page_index'], bad['locator']['page_number'] = 11, 12
    assert (await client.post('/v1/artifact-sources/read', params=params, json={'source': bad})).status_code == 404
    markdown = next(entry for entry in entries if entry['kind'] == 'markdown')
    content_url = root + '/entries/' + markdown['id'] + '/content'
    downloaded = await client.get(content_url, params=params)
    ranged = await client.get(content_url, params=params, headers={'Range': 'bytes=0-7'})
    assert ranged.status_code == 206 and ranged.content == downloaded.content[:8]
    assert ranged.headers['x-content-type-options'] == 'nosniff'
    assert 'attachment;' in ranged.headers['content-disposition']
    invalid_edit = await client.put(f'/v1/artifact-workspaces/{wid}/artifacts/{markdown["id"]}/content', json={
        'operation_id': uuid4().hex, 'path': markdown['path'], 'expected_revision': generation, 'text': 'cannot rewrite bundle'})
    assert invalid_edit.status_code == 404
    copy_body = {'operation_id': uuid4().hex, 'path': 'editable.md'}
    copied = await client.post(root + '/entries/' + markdown['id'] + '/copy', params=params, json=copy_body)
    assert copied.status_code == 200, copied.text
    replay = await client.post(root + '/entries/' + markdown['id'] + '/copy', params=params, json=copy_body)
    assert replay.json()['receipt'] == copied.json()['receipt']
    copy = copied.json()['receipt']
    assert copy['artifact_id'] != ref['artifact_id']
    async with ArtifactsUnitOfWork(engine) as uow:
        row = await uow.session.get(RevisionModel, copy['revision_id'])
        assert row.origin == 'copy' and row.source_representation_set_id == generation
        assert row.source_representation_entry_id == markdown['id']
        count = await uow.session.scalar(select(func.count()).select_from(StoredObject).where(StoredObject.workspace_id == wid))
        ws = await uow.session.get(WorkspaceModel, wid)
        used = ws.used_bytes
        assert ws.reserved_bytes == 0 and used > len(original.encode())
        from modules.artifacts.infra.database.models import ArtifactEventModel
        assert await uow.session.scalar(select(func.count()).select_from(ArtifactEventModel).where(
            ArtifactEventModel.workspace_id == wid, ArtifactEventModel.kind == 'artifact.representations.ready')) == 1
    # A new request uses the exact ready bundle and performs no parsing/output writes.
    response = await client.post(path, params=params, json={**body, 'operation_id': uuid4().hex})
    second = response.json()
    from modules.artifacts.infra.processing.native_text import NativeTextProcessor
    async def unexpected(*args, **kwargs):
        raise AssertionError('ready bundle must avoid repeated parsing')
    monkeypatch.setattr(NativeTextProcessor, 'process', unexpected)
    await build_processing_runner(engine, WorkerSettings.model_validate({})).run_once('bundle-cache-worker')
    second = await completed(client, path, second, params)
    assert second['status'] == 'succeeded' and second['result']['representation_set_id'] == generation, second
    async with ArtifactsUnitOfWork(engine) as uow:
        assert await uow.session.scalar(select(func.count()).select_from(StoredObject).where(StoredObject.workspace_id == wid)) == count
        assert (await uow.session.get(WorkspaceModel, wid)).used_bytes == used
    with pytest.raises(DBAPIError):
        async with ArtifactsUnitOfWork(engine) as uow:
            await uow.session.execute(text('UPDATE artifacts.representation_entries SET path=\'tampered\' WHERE representation_set_id=:id'), {'id': generation})
    # Live source deletion immediately closes historical bundle and SourceRef access.
    async with ArtifactsUnitOfWork(engine) as uow:
        await uow.session.execute(text('UPDATE artifacts.artifacts SET deleted=true WHERE id=:id'), {'id': ref['artifact_id']})
    assert (await client.get(root, params=params)).status_code == 404
    assert (await client.post('/v1/artifact-sources/read', params=params, json={'source': source})).status_code == 404


async def test_quota_failure_keeps_charged_outputs_invisible_and_never_reports_ready(http_client, engine, object_store, monkeypatch):
    client, wid, ref, original, path, body, job, params = await prepared(http_client, engine, object_store, monkeypatch, quota=128)
    await build_processing_runner(engine, WorkerSettings.model_validate({})).run_once('quota-worker')
    status = await completed(client, path, job, params)
    assert status['status'] == 'failed' and status['error_code'] == 'QUOTA_EXCEEDED', status
    async with ArtifactsUnitOfWork(engine) as uow:
        assert await uow.session.scalar(select(func.count()).select_from(Generation).where(Generation.workspace_id == wid, Generation.status == 'ready')) == 0
        ws = await uow.session.get(WorkspaceModel, wid)
        assert ws.used_bytes + ws.reserved_bytes <= ws.quota_bytes


@pytest.mark.parametrize('fault', ['missing_object', 'manifest', 'figure', 'finish_transaction'])
async def test_incomplete_output_and_manifest_never_publish_and_remain_accounted(http_client, engine, object_store, monkeypatch, fault):
    from dataclasses import replace
    from modules.artifacts.domain.models import DocumentBlock, Locator
    from shared.contracts.object_storage import ObjectStorageError
    client, wid, ref, original, path, body, job, params = await prepared(http_client, engine, object_store, monkeypatch)
    if fault == 'missing_object':
        stat = object_store.stat
        async def missing(key):
            if '/markdown-' in key:
                raise ObjectStorageError('not_found', 'Injected missing output')
            return await stat(key)
        monkeypatch.setattr(object_store, 'stat', missing)
    if fault == 'figure':
        from modules.artifacts.infra.processing.native_text import NativeTextProcessor
        process = NativeTextProcessor.process
        async def incomplete(self, *args, **kwargs):
            document = await process(self, *args, **kwargs)
            return replace(document, blocks=document.blocks + (DocumentBlock('missing-figure', 'figure', '', Locator()),))
        monkeypatch.setattr(NativeTextProcessor, 'process', incomplete)
    if fault == 'finish_transaction':
        from modules.artifacts.infra.database.jobs import ProcessingJobRepository
        original_finish = ProcessingJobRepository.finish
        async def broken_finish(self, phase):
            if phase.status == 'succeeded':
                raise RuntimeError('Injected failure after ready mutation, before transaction commit')
            return await original_finish(self, phase)
        monkeypatch.setattr(ProcessingJobRepository, 'finish', broken_finish)
    runner = build_processing_runner(engine, WorkerSettings.model_validate({}))
    if fault == 'manifest':
        phase = runner._publication_phase
        async def inconsistent(name, request):
            if name == 'seal_bundle':
                request = replace(request, value={**request.value, 'metadata_fingerprint': '0' * 64})
            return await phase(name, request)
        runner._publication_phase = inconsistent
    await runner.run_once('incomplete-bundle-worker')
    status = await completed(client, path, job, params)
    assert status['status'] in {'failed', 'queued'}, status
    assert status['result'] is None
    if status['status'] == 'queued':
        cancelled = await client.post(path + '/' + job['id'] + '/cancel', params=params)
        assert cancelled.status_code == 202, cancelled.text
    async with ArtifactsUnitOfWork(engine) as uow:
        assert await uow.session.scalar(select(func.count()).select_from(Generation).where(Generation.workspace_id == wid, Generation.status == 'ready')) == 0
        outputs = (await uow.session.execute(select(StoredObject).where(StoredObject.workspace_id == wid))).scalars().all()
        ws = await uow.session.get(WorkspaceModel, wid)
        assert outputs and ws.used_bytes > len(original.encode())
        assert ws.used_bytes == len(original.encode()) + sum(obj.size_bytes for obj in outputs if obj.accounting_bucket == 'used')
        assert ws.reserved_bytes == sum(obj.size_bytes for obj in outputs if obj.accounting_bucket == 'reserved')
        from modules.artifacts.infra.database.models import ArtifactEventModel
        assert await uow.session.scalar(select(func.count()).select_from(ArtifactEventModel).where(
            ArtifactEventModel.workspace_id == wid, ArtifactEventModel.kind == 'artifact.representations.ready')) == 0


async def test_sealed_bundle_is_invisible_until_current_fence_and_job_commit_together(http_client, engine, object_store, monkeypatch):
    from datetime import UTC, datetime, timedelta
    from modules.artifacts.application.contracts.jobs import ProcessingLeaseLost
    client, wid, ref, original, path, body, job, params = await prepared(http_client, engine, object_store, monkeypatch)
    runner = build_processing_runner(engine, WorkerSettings.model_validate({}))
    finish, captured = runner._finish, []
    async def lose_completion(phase):
        captured.append(phase)
        return True
    runner._finish = lose_completion
    await runner.run_once('sealed-before-crash')
    assert len(captured) == 1 and captured[0].status == 'succeeded'
    generation = captured[0].publication['representation_set_id']
    async with ArtifactsUnitOfWork(engine) as uow:
        root = await uow.session.get(Generation, generation)
        assert root.sealed and root.status == 'building'
        count = await uow.session.scalar(select(func.count()).select_from(StoredObject).where(StoredObject.workspace_id == wid))
        (await uow.session.get(ProcessingJobModel, job['id'])).lease_until = datetime.now(UTC) - timedelta(seconds=1)
    with pytest.raises(ProcessingLeaseLost):
        await finish(captured[0])
    from modules.artifacts.infra.processing.native_text import NativeTextProcessor
    async def unexpected(*args, **kwargs):
        raise AssertionError('sealed bundle recovery must not parse again')
    monkeypatch.setattr(NativeTextProcessor, 'process', unexpected)
    await build_processing_runner(engine, WorkerSettings.model_validate({})).run_once('sealed-after-crash')
    complete = await completed(client, path, job, params)
    assert complete['status'] == 'succeeded' and complete['result']['representation_set_id'] == generation, complete
    async with ArtifactsUnitOfWork(engine) as uow:
        root = await uow.session.get(Generation, generation)
        row = await uow.session.get(ProcessingJobModel, job['id'])
        assert root.status == 'ready' and root.publication_fencing_token == row.fencing_token == 2
        assert await uow.session.scalar(select(func.count()).select_from(StoredObject).where(StoredObject.workspace_id == wid)) == count


async def test_new_generation_and_new_source_revision_cannot_be_replaced_by_an_old_completion(http_client, engine, object_store, monkeypatch):
    client, wid, ref, original, path, body, job, params = await prepared(http_client, engine, object_store, monkeypatch)
    runner = build_processing_runner(engine, WorkerSettings.model_validate({}))
    finish, captured = runner._finish, []
    async def pause(phase):
        captured.append(phase)
        return True
    runner._finish = pause
    await runner.run_once('old-generation-before-completion')
    assert len(captured) == 1 and captured[0].publication
    client_fixture = http_client
    profile = await create_profile(client, client_fixture[4], recipe={'mode': 'extract',
        'processor': {'name': 'native-text', 'version': '1'}, 'parameters': {'preserve_blank_lines': False}})
    assert profile.status_code == 201, profile.text
    newer = await client.post(path, params=params, json={**body, 'operation_id': uuid4().hex, 'profile_id': profile.json()['id']})
    assert newer.status_code == 202, newer.text
    await build_processing_runner(engine, WorkerSettings.model_validate({})).run_once('new-generation-completes-first')
    newer = await completed(client, path, newer.json(), params)
    assert newer['status'] == 'succeeded', newer
    new_generation = newer['result']['representation_set_id']
    # Updating the original while the older result is sealed creates a new
    # immutable source revision, with no inherited processing head.
    updated = await client.put(f"/v1/artifact-workspaces/{wid}/artifacts/{ref['artifact_id']}/content", json={
        'operation_id': uuid4().hex, 'path': 'source.txt', 'text': 'Nova revisão sem processamento.',
        'expected_revision': ref['revision_id']})
    assert updated.status_code == 200, updated.text
    await finish(captured[0])
    base = f"/v1/artifacts/{ref['artifact_id']}/revisions/"
    current = await client.get(base + ref['revision_id'] + '/representations', params=params)
    assert current.status_code == 200 and current.json()['id'] == new_generation, current.text
    new_revision = updated.json()['receipt']['revision_id']
    assert (await client.get(base + new_revision + '/representations', params=params)).status_code == 409
    historical = await client.get(base + ref['revision_id'] + '/representations/' + captured[0].publication['representation_set_id'], params=params)
    assert historical.status_code == 200, historical.text


async def test_cancellation_after_sealing_prevents_ready_and_live_authority_is_required_for_reads(http_client, engine, object_store, monkeypatch):
    from modules.artifacts.application.contracts.jobs import ProcessingCancellationRequested, WorkerPhase
    client, wid, ref, original, path, body, job, params = await prepared(http_client, engine, object_store, monkeypatch)
    runner = build_processing_runner(engine, WorkerSettings.model_validate({}))
    finish, captured = runner._finish, []
    async def pause(phase):
        captured.append(phase)
        return True
    runner._finish = pause
    await runner.run_once('sealed-before-cancel')
    assert captured[0].publication
    response = await client.post(path + '/' + job['id'] + '/cancel', params=params)
    assert response.status_code == 202 and response.json()['status'] == 'cancel_requested', response.text
    with pytest.raises(ProcessingCancellationRequested):
        await finish(captured[0])
    await finish(WorkerPhase(captured[0].lease, status='cancelled'))
    assert (await completed(client, path, job, params))['status'] == 'cancelled'
    async with ArtifactsUnitOfWork(engine) as uow:
        root = await uow.session.get(Generation, captured[0].publication['representation_set_id'])
        assert root.status == 'building'
    new_job = await client.post(path, params=params, json={**body, 'operation_id': uuid4().hex})
    await build_processing_runner(engine, WorkerSettings.model_validate({})).run_once('published-before-revocation')
    complete = await completed(client, path, new_job.json(), params)
    assert complete['status'] == 'succeeded', complete
    base = f"/v1/artifacts/{ref['artifact_id']}/revisions/{ref['revision_id']}/representations"
    response = await client.get(base, params=params)
    assert response.status_code == 200, response.text
    generation = response.json()['id']
    entries = (await client.get(base + '/' + generation + '/entries', params=params)).json()['items']
    entry = next(item for item in entries if item['kind'] == 'source_map')
    downloaded = await client.get(base + '/' + generation + '/entries/' + entry['id'] + '/content', params=params)
    source = downloaded.json()['sources'][0]['source']
    assert (await client.post('/v1/artifact-sources/read', params=params, json={'source': source})).status_code == 200
    async with ArtifactsUnitOfWork(engine) as uow:
        await uow.session.execute(text("UPDATE core.api_keys SET interfaces='[\"http\"]' WHERE id=:id"), {'id': http_client[3].api_key_id})
    assert (await client.get(base, params=params)).status_code == 403
    assert (await client.post('/v1/artifact-sources/read', params=params, json={'source': source})).status_code == 403


@pytest.mark.parametrize('indexed', [False, True])
async def test_real_docling_hierarchy_larger_than_metadata_batch_keeps_parent_foreign_keys_valid(http_client, engine, object_store, monkeypatch, indexed):
    from modules.artifacts.infra.processing.docling_manifest import docling_readiness
    if not docling_readiness()[0]:
        pytest.skip('optional actual Docling converter is unavailable')
    import io
    from docx import Document
    client, wid, ref, original, path, body, job, params = await prepared(http_client, engine, object_store, monkeypatch)
    cancelled = await client.post(path + '/' + job['id'] + '/cancel', params=params)
    assert cancelled.status_code == 202
    document = Document()
    for index in range(105):
        document.add_heading(f'Seção {index} com ação', level=1)
        document.add_paragraph(f'Parágrafo {index}: origem preservada.')
    buffer = io.BytesIO()
    document.save(buffer)
    content = buffer.getvalue()
    media = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    upload_base = f'/v1/artifact-workspaces/{wid}/uploads'
    intent = await client.post(upload_base, json={'operation_id': uuid4().hex, 'path': 'large-hierarchy.docx',
        'size_bytes': len(content), 'sha256': hashlib.sha256(content).hexdigest(), 'content_type': media})
    assert intent.status_code == 201, intent.text
    uploaded = await client.put(upload_base + '/' + intent.json()['id'] + '/parts', params={'offset': 0},
        content=content, headers={'Content-Type': 'application/octet-stream'})
    assert uploaded.status_code == 200, uploaded.text
    published = await client.post(upload_base + '/' + intent.json()['id'] + '/finalize')
    assert published.status_code == 200, published.text
    ref = published.json()['receipt']
    recipe = {'mode': 'extract_index' if indexed else 'extract',
              'processor': {'name': 'docling-native', 'version': '2.140.0'}}
    if indexed:
        recipe['indexer'] = {'name': 'structural', 'version': '1'}
    profile = await create_profile(client, http_client[4], recipe=recipe)
    assert profile.status_code == 201, profile.text
    path = '/v1/artifacts/' + ref['artifact_id'] + '/processing'
    requested = await client.post(path, params=params, json={'operation_id': uuid4().hex,
        'revision_id': ref['revision_id'], 'profile_id': profile.json()['id'], 'profile_revision': 1})
    assert requested.status_code == 202, requested.text
    await build_processing_runner(engine, WorkerSettings.model_validate({})).run_once('large-docling-hierarchy')
    complete = await completed(client, path, requested.json(), params)
    assert complete['status'] == 'succeeded', complete
    assert complete['result']['block_count'] == 210
    if indexed:
        from modules.artifacts.infra.database.index_models import ArtifactDocumentIndexModel
        async with ArtifactsUnitOfWork(engine) as index_uow:
            tree = await index_uow.session.get(ArtifactDocumentIndexModel, complete['result']['document_index_id'])
            assert tree.status == 'ready' and tree.tree_kind == 'canonical' and tree.node_count == 211
    async with ArtifactsUnitOfWork(engine) as uow:
        from modules.artifacts.infra.database.representation_models import ArtifactSourceBlockModel
        rows = (await uow.session.execute(select(ArtifactSourceBlockModel).where(
            ArtifactSourceBlockModel.representation_set_id == complete['result']['representation_set_id']))).scalars().all()
        assert len(rows) == 210 and sum(row.parent_block_id is not None for row in rows) == 105
