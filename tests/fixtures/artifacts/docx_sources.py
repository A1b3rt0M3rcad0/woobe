"""First-party synthetic Office fixture, with no downloaded assets."""
import io

from docx import Document
from docx.shared import Inches
from PIL import Image


def synthetic_docx():
    document = Document()
    document.add_heading('Apostila de ação', level=1)
    document.add_paragraph('Texto português com Unicode: João e maçã.')
    document.add_heading('Tabela e figura', level=2)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = 'Item', 'Valor'
    table.cell(1, 0).text, table.cell(1, 1).text = 'Ação', '12,50'
    document.add_paragraph('Primeiro item', style='List Bullet')
    document.add_paragraph('Segundo item', style='List Bullet')
    image = Image.new('RGB', (64, 32), color=(200, 40, 20))
    picture = io.BytesIO()
    image.save(picture, format='PNG')
    image.close()
    picture.seek(0)
    document.add_picture(picture, width=Inches(1))
    document.add_paragraph('Figura de teste, sem descrição inferida.')
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()
